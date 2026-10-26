"""Domínio financeiro do Horizonte. Valores persistidos em centavos."""
import sqlite3, json, hashlib, re
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from datetime import date, datetime
from pathlib import Path
CATEGORIES=['Moradia e contas','Alimentação','Transporte','Educação','Lazer','Saúde e cuidados','Compras','Assinaturas','Impostos e tarifas','Doações','Empréstimos','Faturas a detalhar','Outros','Salário','Receitas extras','Investimentos']
KINDS=['Receita','Despesa','Aporte','Resgate','Transferência']
MONTHS=['Janeiro','Fevereiro','Março','Abril','Maio','Junho','Julho','Agosto','Setembro','Outubro','Novembro','Dezembro']
def cents(value):
    if isinstance(value,str):
        value=value.replace('R$','').replace(' ','').strip()
        if ',' in value: value=value.replace('.','').replace(',','.')
    try: d=Decimal(str(value))
    except InvalidOperation: raise ValueError('Informe um número, como 125,90.')
    if not d.is_finite() or abs(d)>Decimal('100000000000'): raise ValueError('Valor inválido ou acima do limite.')
    return int((d*100).quantize(Decimal('1'),rounding=ROUND_HALF_UP))
def money(n): return 'R$ '+f'{n/100:,.2f}'.replace(',','_').replace('.',',').replace('_','.')
def month_label(s):
    y,m=map(int,s.split('-')); return f'{MONTHS[m-1]} {y}'
def category(text):
    t=text.lower()
    rules=[('Moradia e contas',['agua','água','casa','internet','vivo','enel','gás','g�s']),('Educação',['faculdade','unicid','japan','curso','dio ']),('Transporte',['transporte','top ','uber','taxi']),('Alimentação',['comida','almo','ifood','assa','doce','lanche','fruta','mercado bom','sorvet','sover','torta','cana mania']),('Assinaturas',['assinatura','telegram','youtube','amazon','rukuten']),('Lazer',['steam','clash','festival','betano','gta','livepix','dorama','quadro']),('Saúde e cuidados',['barbeiro','cabelo','garganta','seguro','perfume']),('Impostos e tarifas',['simples','taxa']),('Doações',['igreja','pratroc']),('Empréstimos',['emprest','sulamita','sula ']),('Faturas a detalhar',['fatura','credito investidor','divida']),('Compras',['shopp','pagleve','pichau','televis','livraria'])]
    for cat,terms in rules:
        if any(x in t for x in terms): return cat
    return 'Outros'
def forecast(initial,contribution,annual,months):
    if not -99<=annual<=1000 or not 0<=months<=600: raise ValueError('Taxa ou prazo fora do intervalo.')
    r=(1+annual/100)**(1/12)-1; amount=float(initial); out=[round(amount)]
    for _ in range(months):
        amount=amount*(1+r)+contribution; out.append(round(amount))
    return out
def required_contribution(initial,target,annual,months):
    if months<=0: return max(0,target-initial)
    r=(1+annual/100)**(1/12)-1
    factor=((1+r)**months-1)/r if abs(r)>1e-12 else months
    return max(0,round((target-initial*(1+r)**months)/factor))
class Store:
    def __init__(self,path):
        self.path=Path(path); self.path.parent.mkdir(parents=True,exist_ok=True)
        self.db=sqlite3.connect(str(path)); self.db.row_factory=sqlite3.Row
        self.db.execute('PRAGMA foreign_keys=ON')
        self.db.executescript('''
        CREATE TABLE IF NOT EXISTS tx(id INTEGER PRIMARY KEY,period TEXT NOT NULL,day TEXT,description TEXT NOT NULL,kind TEXT NOT NULL,category TEXT NOT NULL,amount INTEGER,account TEXT DEFAULT '',status TEXT DEFAULT 'Confirmado',source TEXT DEFAULT '',source_key TEXT UNIQUE,note TEXT DEFAULT '');
        CREATE TABLE IF NOT EXISTS snapshots(id INTEGER PRIMARY KEY,period TEXT NOT NULL,day TEXT,name TEXT,group_name TEXT,amount INTEGER,source TEXT,source_key TEXT UNIQUE);
        CREATE TABLE IF NOT EXISTS positions(id INTEGER PRIMARY KEY,name TEXT NOT NULL,institution TEXT,asset_type TEXT,principal INTEGER,value INTEGER,day TEXT,maturity TEXT,note TEXT);
        CREATE TABLE IF NOT EXISTS budgets(category TEXT,period TEXT,amount INTEGER,PRIMARY KEY(category,period));
        CREATE TABLE IF NOT EXISTS goals(id INTEGER PRIMARY KEY,name TEXT,target INTEGER,saved INTEGER,deadline TEXT,contribution INTEGER,annual REAL);
        CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY,value TEXT);
        CREATE TABLE IF NOT EXISTS imports(hash TEXT PRIMARY KEY,name TEXT,stamp TEXT,report TEXT);
        ''')
        self.db.execute('CREATE TABLE IF NOT EXISTS accounts(id INTEGER PRIMARY KEY,name TEXT UNIQUE NOT NULL,institution TEXT,currency TEXT,balance INTEGER,fx_rate REAL,day TEXT,ledger_after INTEGER,note TEXT)')
        tx_columns={r[1] for r in self.db.execute('PRAGMA table_info(tx)')}
        if 'target_account' not in tx_columns:self.db.execute("ALTER TABLE tx ADD COLUMN target_account TEXT DEFAULT ''")
        columns={r[1] for r in self.db.execute('PRAGMA table_info(positions)')}
        for key,definition in [('asset_group',"TEXT DEFAULT 'A classificar'"),('currency',"TEXT DEFAULT 'BRL'"),('fx_rate','REAL DEFAULT 1'),('native_value','INTEGER')]:
            if key not in columns:self.db.execute('ALTER TABLE positions ADD COLUMN '+key+' '+definition)
        columns={r[1] for r in self.db.execute('PRAGMA table_info(positions)')}
        for key,definition in [('rate_mode',"TEXT DEFAULT 'Não informada'"),('rate','REAL DEFAULT 0'),('tax','REAL DEFAULT 0'),('liquidity',"TEXT DEFAULT 'A confirmar'")]:
            if key not in columns:self.db.execute('ALTER TABLE positions ADD COLUMN '+key+' '+definition)
        if not self.db.execute('SELECT 1 FROM goals').fetchone(): self.db.execute('INSERT INTO goals VALUES (1,?,?,?,?,?,?)',('Meu objetivo',10000000,0,'2030-03-01',0,0))
        self.db.commit()
    def rows(self,sql,args=()): return [dict(r) for r in self.db.execute(sql,args)]
    def setting(self,key,default=''):
        row=self.db.execute('SELECT value FROM settings WHERE key=?',(key,)).fetchone(); return row[0] if row else default
    def set(self,key,value):
        self.db.execute('INSERT OR REPLACE INTO settings VALUES (?,?)',(key,str(value))); self.db.commit()
    def periods(self): return sorted(set([date.today().strftime('%Y-%m')]+[r[0] for r in self.db.execute('SELECT DISTINCT period FROM tx UNION SELECT DISTINCT period FROM snapshots')]))
    def summary(self,period):
        vals={k:0 for k in KINDS}
        for r in self.rows("SELECT kind,SUM(amount) AS total FROM tx WHERE period=? AND amount IS NOT NULL AND status!='Previsto' GROUP BY kind",(period,)): vals[r['kind']]=r['total'] or 0
        vals['livre']=vals['Receita']-vals['Despesa']-vals['Aporte']+vals['Resgate']; return vals
    def save_tx(self,data,ident=None):
        if not data['description'].strip(): raise ValueError('Preencha a descrição.')
        if data['day']: date.fromisoformat(data['day'])
        datetime.strptime(data['period'],'%Y-%m')
        if data['amount'] is not None and data['amount']<0: raise ValueError('Use um valor positivo e selecione o tipo.')
        if data['kind'] not in KINDS: raise ValueError('Tipo inválido.')
        if data['kind']=='Transferência' and data.get('target_account') and data.get('target_account')==data.get('account'):raise ValueError('Origem e destino da transferência devem ser diferentes.')
        cols=['period','day','description','kind','category','amount','account','target_account','status','note']; args=[data.get(k,'') for k in cols]
        if ident: self.db.execute('UPDATE tx SET '+','.join(k+'=?' for k in cols)+' WHERE id=?',args+[ident])
        else: self.db.execute('INSERT INTO tx ('+','.join(cols)+') VALUES ('+','.join('?' for _ in cols)+')',args)
        self.db.commit()
    def backup(self,path):
        if Path(path).resolve()==self.path.resolve(): raise ValueError('Escolha outro arquivo.')
        target=sqlite3.connect(str(path))
        try: self.db.backup(target)
        finally: target.close()
    def account_balance(self,account):
        balance=round(account['balance']*account['fx_rate'])
        for tx in self.rows("SELECT * FROM tx WHERE id>? AND status='Confirmado' AND day>=? AND amount IS NOT NULL AND (account=? OR target_account=?)",(account['ledger_after'],account['day'],account['name'],account['name'])):
            if tx['kind']=='Transferência':
                if tx['account']==account['name']:balance-=tx['amount']
                if tx['target_account']==account['name']:balance+=tx['amount']
            elif tx['account']==account['name']:
                balance+=tx['amount'] if tx['kind'] in ['Receita','Resgate'] else -tx['amount']
        return balance
    def daily_backup(self):
        import os
        folder=Path(self.setting('backup-folder',str(self.path.parent/'backups')))
        folder.mkdir(parents=True,exist_ok=True)
        target=folder/('Horizonte-'+date.today().isoformat()+'.sqlite');temp=folder/(target.name+'.tmp')
        if target.resolve()==self.path.resolve():raise ValueError('Destino de backup inválido.')
        self.backup(temp);os.replace(temp,target);self.set('last-backup',datetime.now().isoformat());return target
    def import_excel(self,path):
        from openpyxl import load_workbook
        digest=hashlib.sha256(Path(path).read_bytes()).hexdigest()
        old=self.rows('SELECT report FROM imports WHERE hash=?',(digest,))
        if old: return 'Arquivo já importado. Nenhuma duplicação.\n'+old[0]['report']
        wb=load_workbook(path,data_only=True,read_only=True); count=snap=blank=0; issues=[]; controls=[]
        try:
            with self.db:
                for sheet in wb:
                    title=sheet.title.upper()
                    if 'FLUXO' in title: continue
                    match=re.search(r'(20\d{2})',title)
                    if not match: continue
                    month=next((i+1 for i,n in enumerate(MONTHS) if title.startswith(n.upper()[:3])),None)
                    if not month: continue
                    period=f'{match[1]}-{month:02d}'; section=''; started=False; local_in=local_out=0; source_total=None
                    for rownum,row in enumerate(sheet.iter_rows(values_only=True),1):
                        a,b,c,d=(list(row)+[None]*6)[:4]; text=str(a or '').strip(); up=text.upper()
                        if up in ['ENTRADAS','ENTRDAS']: section='in'; continue
                        if up in ['SAIDAS','SAÍDAS']: section='out'; continue
                        if up.startswith('TOTAL DE SA'):
                            if isinstance(c,(float,int)): source_total=cents(c)
                            section=''; continue
                        if up.startswith('TOTAL') or up.startswith('SALDO DE ENTRADA'): section=''; continue
                        key=f'{sheet.title}:{rownum}'; source=f'{Path(path).name} / {sheet.title} / linha {rownum}'
                        if section=='in' and isinstance(a,(date,datetime)):
                            desc=' '.join(str(x).strip() for x in [b,c] if x is not None).strip() or 'Entrada sem descrição'
                            if 'salario' in desc.lower() or 'salário' in desc.lower(): started=True
                            amount=cents(d) if isinstance(d,(int,float)) else None
                            day=a.date().isoformat() if isinstance(a,datetime) else a.isoformat()
                            if not started:
                                group='Investimentos a classificar' if (b and any(x in str(b).lower() for x in ['cofrinho','investimento','cdb'])) or any(x in str(c).lower() for x in ['binance','crypto','bitcoin']) else 'Contas e dinheiro'
                                cur=self.db.execute('INSERT OR IGNORE INTO snapshots(period,day,name,group_name,amount,source,source_key) VALUES (?,?,?,?,?,?,?)',(period,day,desc,group,amount,source,key)); snap+=cur.rowcount; continue
                            note='Classificação sugerida. Confirme a natureza da entrada.'
                            if day[:7]!=period: note+=' Data original fora da competência.'; issues.append(source+': data fora do mês')
                            cat='Salário' if 'sal' in desc.lower() else 'Receitas extras'; kind='Receita'; local_in+=amount or 0
                        elif section=='out' and text and not up.startswith('DESCRI'):
                            desc=text; amount=cents(c) if isinstance(c,(int,float)) else None; day=None
                            kind='Aporte' if re.match(r'^(SOFISA\s+\d|Cofrinho PICPAY)',desc,re.I) else 'Despesa'
                            cat='Investimentos' if kind=='Aporte' else category(desc)
                            note='Sem data estruturada na origem. Categoria e natureza sugeridas; revise.'
                            if cat=='Faturas a detalhar': note+=' Fatura agregada: não adicionar compras já incluídas sem ajustar este total.'
                            local_out+=amount or 0
                        else: continue
                        if amount is None: blank+=1; note+=' Valor ausente: fora dos totais.'
                        cur=self.db.execute('INSERT OR IGNORE INTO tx(period,day,description,kind,category,amount,account,status,source,source_key,note) VALUES (?,?,?,?,?,?,?,?,?,?,?)',(period,day,desc,kind,cat,amount,'Origem Excel','A revisar',source,key,note)); count+=cur.rowcount
                    if source_total is not None and local_out!=source_total: issues.append(f'{sheet.title}: saídas diferem do total em {money(local_out-source_total)}')
                    controls.append(f'{month_label(period)}: receitas {money(local_in)}; saídas com aportes {money(local_out)}')
                if count==0 and snap==0: raise ValueError('Layout não reconhecido. Use o layout mensal documentado.')
                report=f'{count} lançamentos e {snap} posições históricas importados.\n{blank} valores ausentes, excluídos dos totais.\n\n'+'\n'.join(controls)+'\n\nRevisões:\n'+'\n'.join(issues)+'\nResumo anual não importado: apresenta diferenças de critério e de totalização. Valores importados não confirmam pagamento. Datas ausentes permanecem sem dia definido. Uma versão alterada do arquivo adiciona somente linhas com novas origens; não sobrescreve linhas já importadas.'
                self.db.execute('INSERT INTO imports VALUES (?,?,?,?)',(digest,Path(path).name,datetime.now().isoformat(),report))
            return report
        finally: wb.close()

